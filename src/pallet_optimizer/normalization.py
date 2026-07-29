from __future__ import annotations

import csv
import io
from collections.abc import Mapping
from typing import Any

from openpyxl import load_workbook

from .catalog import default_vehicle_catalog, find_vehicle
from .domain import CargoItem, Diagnostic, DomainError, Margins, OptimizationProblem, Shape, VehiclePolicy, VehicleVersion


_DIMENSION_FACTORS = {"mm": 1.0, "cm": 10.0, "m": 1000.0}
_WEIGHT_FACTORS = {"kg": 1.0, "g": 0.001, "t": 1000.0}


def _number(value: Any, field: str) -> float:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError) as exc:
        raise DomainError(Diagnostic("INVALID_NUMBER", f"{field} must be numeric", field_path=field)) from exc


def _bool(value: Any, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "oui", "y", "o"}


def _tokens(value: Any) -> tuple[str, ...]:
    if value is None or value == "":
        return ()
    if isinstance(value, (list, tuple)):
        return tuple(str(v).strip() for v in value if str(v).strip())
    return tuple(part.strip() for part in str(value).replace(";", ",").split(",") if part.strip())


def normalize_payload(
    payload: Mapping[str, Any],
    *,
    requested_solutions: int | None = None,
    catalog: tuple[VehicleVersion, ...] | None = None,
) -> OptimizationProblem:
    raw_items = payload.get("items")
    if not isinstance(raw_items, list) or not raw_items:
        raise DomainError(Diagnostic("INVALID_ITEMS", "items must be a non-empty list", field_path="items"))
    default_dimension_unit = str(payload.get("dimension_unit", "mm")).lower()
    default_weight_unit = str(payload.get("weight_unit", "kg")).lower()
    if default_dimension_unit not in _DIMENSION_FACTORS or default_weight_unit not in _WEIGHT_FACTORS:
        raise DomainError(Diagnostic("INVALID_UNIT", "Supported units are mm/cm/m and g/kg/t"))
    defaults = payload.get("default_margins", {}) or {}
    expanded: list[CargoItem] = []
    for input_index, raw in enumerate(raw_items):
        if not isinstance(raw, Mapping):
            raise DomainError(Diagnostic("INVALID_ITEM", f"items[{input_index}] must be an object"))
        dim_unit = str(raw.get("dimension_unit", default_dimension_unit)).lower()
        weight_unit = str(raw.get("weight_unit", default_weight_unit)).lower()
        if dim_unit not in _DIMENSION_FACTORS or weight_unit not in _WEIGHT_FACTORS:
            raise DomainError(Diagnostic("INVALID_UNIT", f"Invalid unit for items[{input_index}]"))
        df, wf = _DIMENSION_FACTORS[dim_unit], _WEIGHT_FACTORS[weight_unit]
        source_id = str(raw.get("id") or f"item-{input_index + 1}")
        quantity = int(_number(raw.get("quantity", 1), f"items[{input_index}].quantity"))
        if quantity < 1:
            raise DomainError(Diagnostic("INVALID_QUANTITY", "quantity must be at least 1", field_path=f"items[{input_index}].quantity"))
        raw_margins = raw.get("margins", {}) or {}
        def margin(name: str) -> int:
            return round(_number(raw_margins.get(name, defaults.get(name, 0)), f"items[{input_index}].margins.{name}") * df)
        margins = Margins(margin("left"), margin("right"), margin("front"), margin("rear"), margin("top"))
        shape_value = str(raw.get("shape", "pallet")).lower()
        try:
            shape = Shape(shape_value)
        except ValueError as exc:
            raise DomainError(Diagnostic("INVALID_SHAPE", f"Unsupported shape: {shape_value}", field_path=f"items[{input_index}].shape")) from exc
        for copy_index in range(quantity):
            item_id = source_id if quantity == 1 else f"{source_id}#{copy_index + 1}"
            expanded.append(CargoItem(
                id=item_id,
                source_id=source_id,
                input_index=input_index,
                shape=shape,
                length_mm=round(_number(raw.get("length"), f"items[{input_index}].length") * df),
                width_mm=round(_number(raw.get("width"), f"items[{input_index}].width") * df),
                height_mm=round(_number(raw.get("height"), f"items[{input_index}].height") * df),
                weight_kg=_number(raw.get("weight"), f"items[{input_index}].weight") * wf,
                destination=str(raw.get("destination") or f"Destination {input_index + 1}"),
                delivery_order=int(_number(raw.get("delivery_order", input_index + 1), f"items[{input_index}].delivery_order")),
                rotation_allowed=_bool(raw.get("rotation_allowed"), True),
                margins=margins,
                compatibility_tags=_tokens(raw.get("compatibility_tags")),
                incompatible_tags=_tokens(raw.get("incompatible_tags")),
                keep_together_group=str(raw["keep_together_group"]) if raw.get("keep_together_group") else None,
                separate_group=str(raw["separate_group"]) if raw.get("separate_group") else None,
                separation_mm=round(_number(raw.get("separation", 0), f"items[{input_index}].separation") * df),
                zone=str(raw["zone"]) if raw.get("zone") else None,
            ))
    catalog = catalog or default_vehicle_catalog()
    policy_raw = payload.get("vehicle_policy", {}) or {}
    mode = str(policy_raw.get("mode", "auto"))
    forced_id = policy_raw.get("forced_vehicle_id")
    policy = VehiclePolicy(mode=mode, forced_vehicle_id=str(forced_id) if forced_id else None,
                           max_vehicles=int(policy_raw.get("max_vehicles", 5)))
    selected = catalog
    if policy.mode == "forced":
        try:
            selected = (find_vehicle(policy.forced_vehicle_id or "", catalog),)
        except KeyError as exc:
            raise DomainError(Diagnostic("UNKNOWN_VEHICLE", f"Unknown vehicle: {policy.forced_vehicle_id}")) from exc
    return OptimizationProblem(
        items=tuple(expanded),
        vehicles=selected,
        vehicle_policy=policy,
        seed=int(payload.get("seed", 1)),
        budget_seconds=min(float(payload.get("budget_seconds", 30)), 30.0),
        requested_solutions=requested_solutions or int(payload.get("requested_solutions", 5)),
    )


COLUMN_ALIASES = {
    "id": {"id", "identifiant", "reference", "référence"},
    "quantity": {"quantity", "quantite", "quantité", "qty"},
    "shape": {"shape", "forme", "type"},
    "length": {"length", "longueur"},
    "width": {"width", "largeur"},
    "height": {"height", "hauteur"},
    "weight": {"weight", "poids"},
    "destination": {"destination", "client"},
    "delivery_order": {"delivery_order", "ordre_livraison", "ordre"},
    "rotation_allowed": {"rotation_allowed", "rotation_autorisee", "rotation_autorisée"},
}


def _canonical_row(row: Mapping[str, Any]) -> dict[str, Any]:
    canonical: dict[str, Any] = {}
    lowered = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    for target, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in lowered:
                canonical[target] = lowered[alias]
                break
    for key, value in lowered.items():
        if key not in canonical:
            canonical.setdefault(key, value)
    return canonical


def payload_from_csv(content: bytes, **base: Any) -> dict[str, Any]:
    text = content.decode("utf-8-sig")
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    rows = [_canonical_row(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]
    return {**base, "items": rows}


def payload_from_xlsx(content: bytes, **base: Any) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {**base, "items": []}
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    items = [_canonical_row(dict(zip(headers, row, strict=False))) for row in rows[1:] if any(v is not None for v in row)]
    return {**base, "items": items}
